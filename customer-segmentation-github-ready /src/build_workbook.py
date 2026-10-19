"""Build the Excel analytical workings with LIVE formulas (workbook/segmentation_workings.xlsx).

The Scoring sheet recomputes the revenue proxy, the Relationship Engagement Index and the segment for every
customer from the cleaned inputs and the Assumptions sheet. src.validate checks the Excel results equal Python.
Run:  python -m src.build_workbook   (then recalculate, e.g. LibreOffice headless, before reading values)
"""
import json
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from .config import ROOT, SCORED, REPORTS, SEGMENTS, SEGMENT_RULES, REVENUE_ASSUMPTIONS as RA, REI_WEIGHTS as W, \
    RECENCY_HALF_LIFE_DAYS, FREQUENCY_CAP, LOGIN_CAP, ADOPTION_CAP, TENURE_CAP_YEARS, VALUE_BANDS, ENGAGEMENT_BANDS
from .features import CARD_REVOLVING_SHARE

OUT = ROOT / "workbook" / "segmentation_workings.xlsx"
FN = "Arial"; BRAND = "0F3B57"
H = Font(name=FN, bold=True, color="FFFFFF", size=10); HF = PatternFill("solid", fgColor=BRAND)
B = Font(name=FN, size=10); BB = Font(name=FN, size=10, bold=True); IN = Font(name=FN, size=10, color="0000FF"); GR = Font(name=FN, size=10, color="008000")
SEGC = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300"]


def head(ws, r, c1, c2):
    for c in range(c1, c2 + 1):
        x = ws.cell(r, c); x.font = H; x.fill = HF; x.alignment = Alignment(wrap_text=True, vertical="center")


def title(ws, t, sub):
    ws["A1"] = t; ws["A1"].font = Font(name=FN, bold=True, size=14, color=BRAND)
    ws["A2"] = sub; ws["A2"].font = Font(name=FN, italic=True, size=9, color="555555")


def val(v):
    if isinstance(v, (float, np.floating)) and np.isnan(v): return None
    return v.item() if isinstance(v, np.generic) else v


def put_df(ws, df, r0=4, fmt=None):
    for j, c in enumerate(df.columns, 1): ws.cell(r0, j, str(c))
    head(ws, r0, 1, len(df.columns))
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            x = ws.cell(i, j, val(v)); x.font = B
            if fmt and df.columns[j - 1] in fmt: x.number_format = fmt[df.columns[j - 1]]
    ws.freeze_panes = ws.cell(r0 + 1, 1)


def main():
    d = pd.read_csv(SCORED); S = json.load(open(REPORTS / "summary.json")); dq = pd.read_csv(REPORTS / "data_quality_log.csv")
    wb = Workbook(); wb.properties.creator = "Ahsan Anas"; wb.properties.title = "Customer Segmentation & Retention – workings (synthetic data)"

    # ---------------- README
    ws = wb.active; ws.title = "README"
    title(ws, "Customer Segmentation & Retention Analytics – analytical workings", "SYNTHETIC DATA · fictional customers · values in USD · author: Ahsan Anas (built with AI assistance)")
    rows = [("Sheet", "Contents"),
            ("Assumptions", "Every model input in one place (blue = editable). Change a value and Scoring + Segment_Summary recalculate."),
            ("Data_Clean", f"{len(d):,} cleaned synthetic customers – the inputs to the scorecard (blank = not available, never imputed)."),
            ("Scoring", "Live formulas: 8 revenue components → est. annual revenue; 5 engagement components → REI; value band × engagement band → segment."),
            ("Segment_Summary", "Live COUNTIFS / SUMIFS / AVERAGEIFS profile of each segment."),
            ("Scenarios", "Hypothetical interventions with formula-driven sizing (assumptions in blue). Illustrative only."),
            ("Pilot", "Holdout pilot sizing with a live minimum-detectable-effect formula (indicative)."),
            ("Sensitivity", "Alternative weights and cut-offs (computed in Python): share of customers keeping the same segment."),
            ("Signals_Model", "Warning-signal lift table, attrition-model deciles and coefficients (Python, synthetic outcome)."),
            ("DQ_Log", "Every cleaning rule: rows affected, action, rationale."),
            ("Data_Dictionary", "Field definitions."), ("", ""),
            ("Labels", "Synthetic observation = what the fictional data shows · Model assumption = a modelling choice · Hypothetical = an intervention scenario, not a result."),
            ("Reconciliation", "src/validate.py confirms Scoring formulas reproduce the Python revenue, REI and segment for every customer.")]
    for i, (a, b) in enumerate(rows, 4):
        ws.cell(i, 1, a).font = BB; ws.cell(i, 2, b).font = B; ws.cell(i, 2).alignment = Alignment(wrap_text=True)
    head(ws, 4, 1, 2); ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 120

    # ---------------- Assumptions
    ws = wb.create_sheet("Assumptions"); title(ws, "Model assumptions (blue = editable inputs)", "All values are illustrative assumptions for a synthetic portfolio")
    A = {}
    items = [("Revenue proxy (annual)", None, None)] + [(k.replace("_", " "), v, "0.00%") for k, v in RA.items()] + \
            [("card revolving share", CARD_REVOLVING_SHARE, "0%"), ("Engagement index weights", None, None)] + \
            [(f"weight: {k}", v, "0") for k, v in W.items()] + \
            [("recency half-life (days)", RECENCY_HALF_LIFE_DAYS, "0"), ("frequency cap (90-day txns)", FREQUENCY_CAP, "0"),
             ("login cap (30-day logins)", LOGIN_CAP, "0"), ("adoption cap (products beyond current account)", ADOPTION_CAP, "0"),
             ("tenure cap (years)", TENURE_CAP_YEARS, "0"), ("Segment bands", None, None),
             ("value band High ≥ (USD)", VALUE_BANDS["High"], "#,##0"), ("value band Mid ≥ (USD)", VALUE_BANDS["Mid"], "#,##0"),
             ("engagement band High ≥ (REI)", ENGAGEMENT_BANDS["High"], "0"), ("engagement band Mid ≥ (REI)", ENGAGEMENT_BANDS["Mid"], "0")]
    r = 4
    for nm, v, f in items:
        if v is None:
            ws.cell(r, 1, nm).font = Font(name=FN, bold=True, color=BRAND, size=11)
        else:
            ws.cell(r, 1, nm).font = B; c = ws.cell(r, 2, v); c.font = IN; c.number_format = f
            A[nm] = f"Assumptions!$B${r}"
        r += 1
    ws.cell(r + 1, 1, "Segment rule table").font = Font(name=FN, bold=True, color=BRAND, size=11)
    ws.cell(r + 2, 1, "Key (value|engagement)"); ws.cell(r + 2, 2, "Segment"); head(ws, r + 2, 1, 2)
    for i, ((vb, eb), sg) in enumerate(SEGMENT_RULES.items(), r + 3):
        ws.cell(i, 1, f"{vb}|{eb}").font = IN; ws.cell(i, 2, sg).font = IN
    RULE = (f"Assumptions!$A${r+3}:$A${r+2+len(SEGMENT_RULES)}", f"Assumptions!$B${r+3}:$B${r+2+len(SEGMENT_RULES)}")
    ws.cell(4, 4, "Sum of weights (must be 100)").font = BB
    ws.cell(5, 4, f"=SUM({A['weight: recency']},{A['weight: frequency']},{A['weight: digital']},{A['weight: adoption']},{A['weight: loyalty']})")
    ws.column_dimensions["A"].width = 48; ws.column_dimensions["B"].width = 26; ws.column_dimensions["D"].width = 30

    # ---------------- Data_Clean
    cols = ["customer_id", "region", "age", "main_bank_flag", "tenure_months", "num_products", "current_account_avg_bal_3m",
            "savings_balance", "term_deposit_balance", "investment_balance", "mortgage_balance", "personal_loan_balance",
            "credit_card_limit", "card_utilisation_pct", "card_spend_90d", "days_since_txn", "txn_count_90d", "app_logins_30d",
            "digital_txn_share", "complaints_12m", "attrited_next_6m", "attrition_risk_model", "dq_flags"]
    ws = wb.create_sheet("Data_Clean"); title(ws, f"Cleaned synthetic customers ({len(d):,})", "Blank = not available in the source; nothing imputed")
    put_df(ws, d[cols], fmt={c: "#,##0.00" for c in cols[6:15]})
    for j in range(1, len(cols) + 1): ws.column_dimensions[L(j)].width = 13
    C = {c: L(j) for j, c in enumerate(cols, 1)}
    N = len(d); DC = lambda c, i: f"Data_Clean!{C[c]}{i}"

    # ---------------- Scoring (live)
    ws = wb.create_sheet("Scoring"); title(ws, "Scoring – live formulas", "Green = pulled from Data_Clean · black = calculated from Assumptions")
    hd = ["customer_id", "rev_current_account", "rev_savings", "rev_term_deposit", "rev_investments", "rev_mortgage",
          "rev_personal_loan", "rev_card_revolve", "rev_card_interchange", "est_annual_revenue", "rei_recency", "rei_frequency",
          "rei_digital", "rei_adoption", "rei_loyalty", "rei", "value_band", "engagement_band", "segment", "expected_revenue_at_risk"]
    for j, h in enumerate(hd, 1): ws.cell(4, j, h)
    head(ws, 4, 1, len(hd))
    K = {h: L(j) for j, h in enumerate(hd, 1)}
    a = A
    for k in range(N):
        i = 5 + k; s = i  # same row number in Data_Clean (both start at row 5)
        x = lambda c: DC(c, s)
        f = {
            "customer_id": f"={x('customer_id')}",
            "rev_current_account": f"={x('current_account_avg_bal_3m')}*{a['current account margin']}",
            "rev_savings": f"=IF({x('savings_balance')}=\"\",0,{x('savings_balance')})*{a['savings margin']}",
            "rev_term_deposit": f"={x('term_deposit_balance')}*{a['term deposit margin']}",
            "rev_investments": f"={x('investment_balance')}*{a['investment fee']}",
            "rev_mortgage": f"={x('mortgage_balance')}*{a['mortgage margin']}",
            "rev_personal_loan": f"={x('personal_loan_balance')}*{a['personal loan margin']}",
            "rev_card_revolve": f"=IF(OR({x('credit_card_limit')}=\"\",{x('card_utilisation_pct')}=\"\"),0,{x('credit_card_limit')}*{x('card_utilisation_pct')}/100)*{a['card revolving share']}*{a['card revolve margin']}",
            "rev_card_interchange": f"={x('card_spend_90d')}*4*{a['card interchange']}",
            "est_annual_revenue": f"=SUM({K['rev_current_account']}{i}:{K['rev_card_interchange']}{i})",
            "rei_recency": f"=IF({x('days_since_txn')}=\"\",\"\",{a['weight: recency']}*0.5^({x('days_since_txn')}/{a['recency half-life (days)']}))",
            "rei_frequency": f"={a['weight: frequency']}*MIN(1,LN(1+{x('txn_count_90d')})/LN(1+{a['frequency cap (90-day txns)']}))",
            "rei_digital": f"={a['weight: digital']}*(0.6*MIN(1,{x('app_logins_30d')}/{a['login cap (30-day logins)']})+0.4*IF({x('digital_txn_share')}=\"\",0,{x('digital_txn_share')}))",
            "rei_adoption": f"={a['weight: adoption']}*MIN(1,({x('num_products')}-1)/{a['adoption cap (products beyond current account)']})",
            "rei_loyalty": (f"={a['weight: loyalty']}*(0.4*IF({x('tenure_months')}=\"\",0,MIN(1,{x('tenure_months')}/12/{a['tenure cap (years)']}))"
                            f"+0.35*IF({x('main_bank_flag')}=\"Y\",1,0)+0.25*(1-MIN({x('complaints_12m')},3)/3))/IF({x('tenure_months')}=\"\",0.6,1)"),
            "rei": f"=SUM({K['rei_recency']}{i}:{K['rei_loyalty']}{i})/(Assumptions!$D$5-IF({K['rei_recency']}{i}=\"\",{a['weight: recency']},0))*100",
            "value_band": f"=IF({K['est_annual_revenue']}{i}>={a['value band High ≥ (USD)']},\"High\",IF({K['est_annual_revenue']}{i}>={a['value band Mid ≥ (USD)']},\"Mid\",\"Low\"))",
            "engagement_band": f"=IF({K['rei']}{i}>={a['engagement band High ≥ (REI)']},\"High\",IF({K['rei']}{i}>={a['engagement band Mid ≥ (REI)']},\"Mid\",\"Low\"))",
            "segment": f"=INDEX({RULE[1]},MATCH({K['value_band']}{i}&\"|\"&{K['engagement_band']}{i},{RULE[0]},0))",
            "expected_revenue_at_risk": f"=IF(ISNUMBER({x('attrition_risk_model')}),{K['est_annual_revenue']}{i}*{x('attrition_risk_model')},0)",
        }
        for h, fx in f.items():
            c = ws[f"{K[h]}{i}"]; c.value = fx; c.font = GR if h == "customer_id" else B
            if h.startswith("rev") or h == "est_annual_revenue": c.number_format = "#,##0.00"
            if h.startswith("rei"): c.number_format = "0.00"
    for j in range(1, len(hd) + 1): ws.column_dimensions[L(j)].width = 13
    ws.column_dimensions[K["segment"]].width = 22; ws.freeze_panes = "B5"
    for sg, cc in zip(SEGMENTS, SEGC):
        ws.conditional_formatting.add(f"{K['segment']}5:{K['segment']}{4+N}", CellIsRule(operator="equal", formula=[f'"{sg}"'],
                                      fill=PatternFill("solid", fgColor=cc), font=Font(color="FFFFFF", bold=True)))
    rng = lambda h: f"Scoring!${K[h]}$5:${K[h]}${4+N}"
    dcr = lambda c: f"Data_Clean!${C[c]}$5:${C[c]}${4+N}"

    # ---------------- Segment_Summary (live)
    ws = wb.create_sheet("Segment_Summary"); title(ws, "Segment summary – live formulas (synthetic observation)", "Recalculates when Assumptions change")
    hd2 = ["Segment", "Customers", "% of customers", "Est. revenue (USD)", "% of revenue", "Avg revenue / customer", "Avg REI",
           "Avg products", "% main bank", "Attrition next 6m (synthetic)", "Expected revenue at risk (model)"]
    for j, h in enumerate(hd2, 1): ws.cell(4, j, h)
    head(ws, 4, 1, len(hd2))
    for i, sg in enumerate(SEGMENTS, 5):
        ws.cell(i, 1, sg).font = Font(name=FN, bold=True, color="FFFFFF"); ws.cell(i, 1).fill = PatternFill("solid", fgColor=SEGC[i - 5])
        sgc = f"$A{i}"
        fx = [f"=COUNTIF({rng('segment')},{sgc})", f"=B{i}/SUM($B$5:$B$10)", f"=SUMIF({rng('segment')},{sgc},{rng('est_annual_revenue')})",
              f"=D{i}/SUM($D$5:$D$10)", f"=D{i}/B{i}", f"=AVERAGEIF({rng('segment')},{sgc},{rng('rei')})",
              f"=AVERAGEIF({rng('segment')},{sgc},{dcr('num_products')})", f"=COUNTIFS({rng('segment')},{sgc},{dcr('main_bank_flag')},\"Y\")/B{i}",
              f"=AVERAGEIF({rng('segment')},{sgc},{dcr('attrited_next_6m')})",
              f"=SUMIF({rng('segment')},{sgc},{rng('expected_revenue_at_risk')})"]
        for j, x_ in enumerate(fx, 2):
            c = ws.cell(i, j, x_); c.font = B
            c.number_format = {3: "0.0%", 4: "#,##0", 5: "0.0%", 6: "#,##0", 7: "0.0", 8: "0.00", 9: "0.0%", 10: "0.0%", 11: "#,##0"}.get(j, "#,##0")
    ws.cell(11, 1, "Total").font = BB
    for j in (2, 3, 4, 5, 11): ws.cell(11, j, f"=SUM({L(j)}5:{L(j)}10)").font = BB
    ws.cell(11, 3).number_format = "0.0%"; ws.cell(11, 5).number_format = "0.0%"; ws.cell(11, 4).number_format = "#,##0"; ws.cell(11, 11).number_format = "#,##0"
    for j, w in enumerate([22, 11, 12, 16, 12, 14, 9, 10, 11, 16, 18], 1): ws.column_dimensions[L(j)].width = w
    ws.row_dimensions[4].height = 42

    # ---------------- Scenarios (live, hypothetical)
    ws = wb.create_sheet("Scenarios"); title(ws, "Hypothetical interventions – illustrative sizing (NOT results)", "Blue = assumption. Outcomes depend entirely on these assumptions.")
    for j, h in enumerate(["Scenario", "Assumption", "Value", "Output", "Result"], 1): ws.cell(4, j, h)
    head(ws, 4, 1, 5)
    ws["A5"] = "Affluent Drifters – RM re-engagement"; ws["B5"] = "Relative cut in attrition probability"; ws["C5"] = 0.25; ws["C5"].font = IN; ws["C5"].number_format = "0%"
    ws["D5"] = "Expected annual revenue protected (USD)"; ws["E5"] = "=Segment_Summary!K7*C5"; ws["E5"].number_format = "#,##0"
    ws["A6"] = "Engaged Builders – cross-sell one product"; ws["B6"] = "Take-up rate"; ws["C6"] = 0.10; ws["C6"].font = IN; ws["C6"].number_format = "0%"
    ws["B7"] = "Annual revenue per added product (USD)"; ws["C7"] = 180; ws["C7"].font = IN
    ws["D6"] = "Incremental annual revenue (USD)"; ws["E6"] = "=Segment_Summary!B8*C6*C7"; ws["E6"].number_format = "#,##0"
    for c in ("A5", "A6"): ws[c].font = BB
    ws["A9"] = "Python reference values: " + "; ".join(f"{s_['segment']}: {s_['value']:,}" for s_ in S["scenarios"])
    ws["A9"].font = Font(name=FN, italic=True, size=9)
    for j, w in enumerate([40, 40, 10, 38, 14], 1): ws.column_dimensions[L(j)].width = w

    # ---------------- Pilot (live MDE)
    ws = wb.create_sheet("Pilot"); title(ws, "Holdout pilot sizing (indicative)", "Minimum detectable change = Z × √(2·p·(1−p)/n) · Z = 2.80 for 95% confidence and 80% power")
    for j, h in enumerate(["Cell", "Eligible", "Treatment", "Holdout (n)", "Baseline rate (p)", "Min. detectable change (pts)"], 1): ws.cell(4, j, h)
    head(ws, 4, 1, 6); ws["H4"] = "Z"; ws["I4"] = 2.80; ws["I4"].font = IN
    for i, c in enumerate(S["pilot"], 5):
        for j, v in enumerate([c["cell"], c["eligible"], c["treatment"], c["holdout"]], 1): ws.cell(i, j, v).font = B
        ws.cell(i, 5, c["baseline_rate_pct"] / 100).font = IN; ws.cell(i, 5).number_format = "0.0%"
        ws.cell(i, 6, f"=$I$4*SQRT(2*E{i}*(1-E{i})/D{i})*100").number_format = "0.0"
    for j, w in enumerate([34, 10, 10, 11, 14, 22], 1): ws.column_dimensions[L(j)].width = w

    # ---------------- static Python tables
    ws = wb.create_sheet("Sensitivity"); title(ws, "Segmentation sensitivity (model assumption tests)", "Computed by src/segment.py")
    put_df(ws, pd.DataFrame(S["sensitivity"])); ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 34
    r = 6 + len(S["sensitivity"])
    ws.cell(r, 1, f"k-means check: silhouette by k {S['kmeans']['silhouette_by_k']} · ARI rule-based vs k-means(6) = {S['kmeans']['ari_rule_vs_kmeans6']}").font = Font(name=FN, italic=True, size=9)
    ws = wb.create_sheet("Signals_Model"); title(ws, "Warning signals and attrition model (synthetic outcome)", f"Logistic regression · test AUC {S['model']['auc_test']} · top-decile lift {S['model']['top_decile_lift']}")
    sg = pd.DataFrame(S["signals"]); put_df(ws, sg)
    r = 6 + len(sg); dec = pd.DataFrame(S["deciles"]); put_df(ws, dec, r0=r)
    r2 = r + 2 + len(dec); put_df(ws, pd.DataFrame(S["coefficients"]), r0=r2)
    ws.freeze_panes = None; ws.column_dimensions["A"].width = 36
    ws = wb.create_sheet("DQ_Log"); title(ws, "Data-quality log", f"{S['rows_raw']:,} raw rows → {S['customers']:,} customers")
    put_df(ws, dq)
    for j, w in enumerate([6, 34, 46, 10, 50, 70, 30], 1): ws.column_dimensions[L(j)].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    dd = pd.read_csv(ROOT / "docs" / "data_dictionary.csv") if (ROOT / "docs" / "data_dictionary.csv").exists() else pd.DataFrame()
    ws = wb.create_sheet("Data_Dictionary"); title(ws, "Data dictionary", "Synthetic fields")
    if len(dd): put_df(ws, dd); ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 90
    wb.calculation.fullCalcOnLoad = True   # Excel recalculates every formula when the file is opened
    OUT.parent.mkdir(exist_ok=True); wb.save(OUT); print("saved", OUT)


if __name__ == "__main__":
    main()
